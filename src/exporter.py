from pathlib import Path
from datetime import datetime

import pandas as pd

from openpyxl import load_workbook
from openpyxl.styles import Font, PatternFill, Alignment
from openpyxl.formatting.rule import ColorScaleRule
from openpyxl.worksheet.table import Table, TableStyleInfo
from openpyxl.chart import BarChart, Reference


class ExcelExporter:

    def __init__(self, output_folder):
        self.output_folder = Path(output_folder)

        self.output_folder.mkdir(
            parents=True,
            exist_ok=True
        )

    def export_report(self, analyzer):
        """
        Export sales analysis results to a formatted Excel workbook.
        """

        output_file = self.output_folder / "sales_report.xlsx"

        # ---------------------------------
        # Create analysis data
        # ---------------------------------

        summary = pd.DataFrame({
            "Metric": [
                "Total Sales",
                "Total Profit",
                "Total Quantity",
                "Average Order Value",
                "Profit Margin"
            ],
            "Value": [
                analyzer.total_sales(),
                analyzer.total_profit(),
                analyzer.total_quantity(),
                analyzer.average_order_value(),
                analyzer.profit_margin() / 100
            ]
        })

        category_sales = analyzer.sales_by_category()
        category_profit = analyzer.profit_by_category()

        category = pd.concat(
        [
        category_sales.rename("Sales"),
        category_profit.rename("Profit")
        ],
        axis=1
        ).reset_index()
        category.columns = ["Category" ,"Sales", "Profit"]

        region_sales = analyzer.sales_by_region()
        region_profit = analyzer.profit_by_region()

        region = pd.concat(
        [
            region_sales.rename("Sales"),
            region_profit.rename("Profit")
        ],
        axis=1
        ).reset_index()
        region.columns = ["Region", "Sales", "Profit"]

        segment = pd.DataFrame({
            "Segment": analyzer.sales_by_segment().index,
            "Sales": analyzer.sales_by_segment().values
        })

        yearly_sales = analyzer.sales_by_year()
        yearly_profit = analyzer.profit_by_year()
        yearly_growth = analyzer.sales_growth_by_year()

        yearly = pd.concat(
            [
                yearly_sales.rename("Sales"),
                yearly_profit.rename("Profit"),
                yearly_growth.rename("Growth")
            ],
            axis=1
        ).reset_index()
        yearly.columns = ["Year", "Sales", "Profit", "Growth"]

        discount = analyzer.profit_by_discount().reset_index()
        discount.columns = ["Discount", "Profit","Records","Average Profit"]

        products = analyzer.top_products().reset_index()
        products.columns = ["Product", "Sales"]

        customers = analyzer.top_customers().reset_index()
        customers.columns = ["Customer", "Sales"]

        # ---------------------------------
        # Create workbook
        # ---------------------------------

        with pd.ExcelWriter(
            output_file,
            engine="openpyxl"
        ) as writer:

            summary.to_excel(
                writer,
                sheet_name="Summary",
                index=False
            )

            category.to_excel(
                writer,
                sheet_name="Category Analysis",
                index=False
            )

            region.to_excel(
                writer,
                sheet_name="Region Analysis",
                index=False
            )

            segment.to_excel(
                writer,
                sheet_name="Segment Analysis",
                index=False
            )
            yearly.to_excel(
                writer,
                sheet_name="Yearly Analysis",
                index=False
            )

            products.to_excel(
                writer,
                sheet_name="Top Products",
                index=False
            )
            discount.to_excel(
                writer,
                sheet_name="Discount Analysis",
                index=False
            )
            
            customers.to_excel(
                writer,
                sheet_name="Top Customers",
                index=False
            )

        # ---------------------------------
        # Format workbook
        # ---------------------------------

        self.format_workbook(output_file)

        # ---------------------------------
        # Create dashboard
        # ---------------------------------

        self.create_dashboard(
            output_file,
            analyzer
        )

        print(
            f"\nProfessional Excel report created successfully:"
            f"\n{output_file}"
        )

        return output_file

    # =================================
    # WORKBOOK FORMATTING
    # =================================

    def format_workbook(self, output_file):

        workbook = load_workbook(output_file)

        header_fill = PatternFill(
            "solid",
            fgColor="1F4E78"
        )

        header_font = Font(
            color="FFFFFF",
            bold=True
        )

        title_font = Font(
            size=16,
            bold=True
        )

        for sheet in workbook.worksheets:

            sheet.freeze_panes = "A2"

            for cell in sheet[1]:

                cell.fill = header_fill
                cell.font = header_font
                cell.alignment = Alignment(
                    horizontal="center"
                )

            sheet.auto_filter.ref = sheet.dimensions

            # Adjust column widths
            for column in sheet.columns:

                max_length = 0
                column_letter = column[0].column_letter

                for cell in column:

                    if cell.value is not None:

                        max_length = max(
                            max_length,
                            len(str(cell.value))
                        )

                sheet.column_dimensions[
                    column_letter
                ].width = min(
                    max_length + 3,
                    45
                )

        # ---------------------------------
        # Summary
        # ---------------------------------

        summary = workbook["Summary"]

        summary.insert_rows(1)

        summary["A1"] = "Retail Sales Performance Report"

        summary["A1"].font = title_font

        summary["A2"] = "Generated:"
        summary["B2"] = datetime.now().strftime(
            "%Y-%m-%d %H:%M"
        )

        summary["A2"].font = Font(
            bold=True
        )

        for cell in summary[4]:

            cell.fill = header_fill
            cell.font = header_font
            cell.alignment = Alignment(
                horizontal="center"
            )

        summary.freeze_panes = "A5"

        summary["B5"].number_format = '#,##0.00'
        summary["B6"].number_format = '#,##0.00'
        summary["B7"].number_format = '#,##0'
        summary["B8"].number_format = '#,##0.00'
        summary["B9"].number_format = '0.00%'

        # ---------------------------------
        # Category & Region
        # ---------------------------------

        for sheet_name in [
            "Category Analysis",
            "Region Analysis"
        ]:

            sheet = workbook[sheet_name]

            for row in range(
                2,
                sheet.max_row + 1
            ):

                sheet[f"B{row}"].number_format = '#,##0.00'
                sheet[f"C{row}"].number_format = '#,##0.00'

            sheet.conditional_formatting.add(
                f"C2:C{sheet.max_row}",
                ColorScaleRule(
                    start_type="min",
                    start_color="F8696B",
                    mid_type="percentile",
                    mid_value=50,
                    mid_color="FFEB84",
                    end_type="max",
                    end_color="63BE7B"
                )
            )
               
        # ---------------------------------
        # Yearly Analysis
        # ---------------------------------

        yearly = workbook["Yearly Analysis"]

        for row in range(
             2,
             yearly.max_row + 1
        ):

             yearly[f"B{row}"].number_format = '#,##0.00'
             yearly[f"C{row}"].number_format = '#,##0.00'
             yearly[f"D{row}"].number_format = '#,##0'


        # ---------------------------------
        # Discount Analysis
        # ---------------------------------

        discount = workbook["Discount Analysis"]

        for row in range(
             2,
             discount.max_row + 1
        ):

            discount[f"B{row}"].number_format = '#,##0.00'
            discount[f"C{row}"].number_format = '#,##0.00'
            discount[f"D{row}"].number_format = '#,##0'    

        # ---------------------------------
        # Segment
        # ---------------------------------

        segment = workbook["Segment Analysis"]

        for row in range(
            2,
            segment.max_row + 1
        ):

            segment[f"B{row}"].number_format = '#,##0.00'

        # ---------------------------------
        # Products
        # ---------------------------------

        products = workbook["Top Products"]

        for row in range(
            2,
            products.max_row + 1
        ):

            products[f"B{row}"].number_format = '#,##0.00'

        # ---------------------------------
        # Customers
        # ---------------------------------

        customers = workbook["Top Customers"]

        for row in range(
            2,
            customers.max_row + 1
        ):

            customers[f"B{row}"].number_format = '#,##0.00'

        # ---------------------------------
        # Tables
        # ---------------------------------

        self.add_table(
            workbook["Category Analysis"],
            "CategoryTable"
        )

        self.add_table(
            workbook["Region Analysis"],
            "RegionTable"
        )

        self.add_table(
            workbook["Segment Analysis"],
            "SegmentTable"
        )

        self.add_table(
            workbook["Top Products"],
            "ProductsTable"
        )

        self.add_table(
            workbook["Top Customers"],
            "CustomersTable"
        )

        self.add_table(workbook["Yearly Analysis"], "YearlyTable")
        self.add_table(workbook["Discount Analysis"], "DiscountTable")

        workbook.save(output_file)

    # =================================
    # EXCEL TABLE
    # =================================

    def add_table(self, sheet, table_name):

        table = Table(
            displayName=table_name,
            ref=sheet.dimensions
        )

        style = TableStyleInfo(
            name="TableStyleMedium2",
            showFirstColumn=False,
            showLastColumn=False,
            showRowStripes=True,
            showColumnStripes=False
        )

        table.tableStyleInfo = style

        sheet.add_table(table)

    # =================================
    # DASHBOARD
    # =================================

    def create_dashboard(
        self,
        output_file,
        analyzer
    ):
        """
        Create an automated Excel dashboard.
        """

        workbook = load_workbook(output_file)

        # Remove old dashboard if it exists
        if "Dashboard" in workbook.sheetnames:

            del workbook["Dashboard"]

        dashboard = workbook.create_sheet(
            "Dashboard",
            0
        )

        # ---------------------------------
        # Title
        # ---------------------------------

        dashboard["B2"] = "RETAIL SALES PERFORMANCE"

        dashboard["B2"].font = Font(
            size=20,
            bold=True
        )

        dashboard["B3"] = (
            "Automated Sales Analysis Dashboard"
        )

        dashboard["B3"].font = Font(
            italic=True,
            size=11
        )

        dashboard["B4"] = (
            "Generated: "
            + datetime.now().strftime(
                "%Y-%m-%d %H:%M"
            )
        )

        # ---------------------------------
        # KPI section
        # ---------------------------------

        kpis = [
            ("B6", "TOTAL SALES", analyzer.total_sales()),
            ("E6", "TOTAL PROFIT", analyzer.total_profit()),
            ("H6", "PROFIT MARGIN", analyzer.profit_margin() / 100)
        ]

        for position, title, value in kpis:

            dashboard[position] = title

            dashboard[position].font = Font(
                bold=True,
                size=11
            )

            value_cell = dashboard.cell(
                row=7,
                column=dashboard[position].column
            )

            value_cell.value = value

            value_cell.font = Font(
                bold=True,
                size=16
            )

            value_cell.alignment = Alignment(
                horizontal="center"
            )

            dashboard.merge_cells(
                start_row=6,
                start_column=dashboard[position].column,
                end_row=6,
                end_column=dashboard[position].column + 1
            )

            dashboard.merge_cells(
                start_row=7,
                start_column=dashboard[position].column,
                end_row=8,
                end_column=dashboard[position].column + 1
            )

        dashboard["B7"].number_format = '#,##0.00'
        dashboard["E7"].number_format = '#,##0.00'
        dashboard["H7"].number_format = '0.00%'

        # ---------------------------------
        # Category chart
        # ---------------------------------

        category_sheet = workbook["Category Analysis"]

        category_chart = BarChart()

        category_chart.type = "bar"
        category_chart.style = 10
        category_chart.title = "Sales by Category"
        category_chart.y_axis.title = "Category"
        category_chart.x_axis.title = "Sales"

        data = Reference(
            category_sheet,
            min_col=2,
            min_row=1,
            max_row=category_sheet.max_row
        )

        categories = Reference(
            category_sheet,
            min_col=1,
            min_row=2,
            max_row=category_sheet.max_row
        )

        category_chart.add_data(
            data,
            titles_from_data=True
        )

        category_chart.set_categories(
            categories
        )

        category_chart.height = 7
        category_chart.width = 11

        dashboard.add_chart(
            category_chart,
            "B11"
        )
        # ---------------------------------
        # Yearly Sales Chart
        # ---------------------------------

        yearly_sheet = workbook["Yearly Analysis"]

        yearly_chart = BarChart()

        yearly_chart.type = "col"
        yearly_chart.style = 10
        yearly_chart.title = "Sales by Year"
        yearly_chart.y_axis.title = "Sales"
        yearly_chart.x_axis.title = "Year"

        data = Reference(
            yearly_sheet,
            min_col=2,
            min_row=1,
            max_row=yearly_sheet.max_row
        )

        years = Reference(
            yearly_sheet,
            min_col=1,
            min_row=2,
            max_row=yearly_sheet.max_row
        )

        yearly_chart.add_data(
            data,
            titles_from_data=True
        )

        yearly_chart.set_categories(
            years
        )

        yearly_chart.height = 7
        yearly_chart.width = 11

        dashboard.add_chart(
            yearly_chart,
            "B26"
        )

        # ---------------------------------
        # Profit by Discount Band Chart
        # ---------------------------------

        discount_sheet = workbook["Discount Analysis"]

        discount_chart = BarChart()

        discount_chart.type = "col"
        discount_chart.style = 11
        discount_chart.title = "Profit by Discount Band"
        discount_chart.y_axis.title = "Profit"
        discount_chart.x_axis.title = "Discount Band"

        data = Reference(
            discount_sheet,
            min_col=2,
            min_row=1,
            max_row=discount_sheet.max_row
        )

        discount_bands = Reference(
            discount_sheet,
            min_col=1,
            min_row=2,
            max_row=discount_sheet.max_row
        )

        discount_chart.add_data(
            data,
            titles_from_data=True
        )

        discount_chart.set_categories(
            discount_bands
        )

        discount_chart.height = 7
        discount_chart.width = 11

        dashboard.add_chart(
            discount_chart,
            "H26"
        )

        # ---------------------------------
        # Region chart
        # ---------------------------------

        region_sheet = workbook["Region Analysis"]

        region_chart = BarChart()

        region_chart.type = "bar"
        region_chart.style = 11
        region_chart.title = "Sales by Region"
        region_chart.y_axis.title = "Region"
        region_chart.x_axis.title = "Sales"

        data = Reference(
            region_sheet,
            min_col=2,
            min_row=1,
            max_row=region_sheet.max_row
        )

        categories = Reference(
            region_sheet,
            min_col=1,
            min_row=2,
            max_row=region_sheet.max_row
        )

        region_chart.add_data(
            data,
            titles_from_data=True
        )

        region_chart.set_categories(
            categories
        )

        region_chart.height = 7
        region_chart.width = 11

        dashboard.add_chart(
            region_chart,
            "H11"
        )

        # ---------------------------------
        # Column widths
        # ---------------------------------

        for column in range(
            1,
            12
        ):

            dashboard.column_dimensions[
                chr(64 + column)
            ].width = 14

        # ---------------------------------
        # Page setup
        # ---------------------------------

        dashboard.sheet_view.showGridLines = False

        dashboard.freeze_panes = "B6"

        workbook.save(output_file)
